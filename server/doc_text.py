"""Şartname dosyalarından metin çıkarma (karar 17): düz metin ve Word (.docx, python-docx).

Görsel/mockup dosyaları (png, jpg, …) kasaya kaydedilir ama içeriği OKUNMAZ (OCR yoktur); `is_unreadable_image` ile ayırt edilir.
"""
from __future__ import annotations

import io
import os
from typing import Optional, Union

TEXT_EXT = (".md", ".txt", ".json", ".markdown")
WORD_EXT = (".docx",)
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg")


def is_readable(name: str) -> bool:
    return str(name).lower().endswith(TEXT_EXT + WORD_EXT)


def is_unreadable_image(name: str) -> bool:
    return str(name).lower().endswith(IMAGE_EXT)


def docx_text(source: Union[str, bytes]) -> str:
    """Paragraflar (başlıklar `#` ile) ve tablolar (satır başına `a | b | c`); bozuk dosyada boş döner."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source)
    except Exception:
        return ""
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            lvl = "".join(c for c in style if c.isdigit()) or "1"
            t = "#" * max(1, min(int(lvl), 4)) + " " + t
        lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_text(path_or_bytes: Union[str, bytes], filename: Optional[str] = None) -> str:
    """Dosya yolundan ya da baytlardan (ad verilerek) okunabilir metin; desteklenmeyen türde boş döner."""
    name = (filename or (path_or_bytes if isinstance(path_or_bytes, str) else "")).lower()
    if name.endswith(WORD_EXT):
        return docx_text(path_or_bytes)
    if name.endswith(TEXT_EXT):
        if isinstance(path_or_bytes, (bytes, bytearray)):
            return bytes(path_or_bytes).decode("utf-8", errors="ignore")
        with open(path_or_bytes, "r", encoding="utf-8", errors="ignore") as fh:
            return fh.read()
    return ""
