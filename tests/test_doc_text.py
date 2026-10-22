"""Karar 17: Word (.docx) şartnameleri okunur; görseller kasaya girer ama okunmaz."""
import io

from server import doc_text as D


def _docx_bytes():
    from docx import Document
    d = Document()
    d.add_heading("Müşteri Portalı", level=1)
    d.add_paragraph("Kullanıcılar sisteme kayıt olup giriş yapabilmeli.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Modül", "Not"
    t.cell(1, 0).text, t.cell(1, 1).text = "Raporlar", "PDF dışa aktarım"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_text_includes_headings_paragraphs_and_tables():
    text = D.read_text(_docx_bytes(), "sartname.DOCX")
    assert "# Müşteri Portalı" in text and "Kullanıcılar sisteme kayıt olup giriş yapabilmeli." in text and "Raporlar | PDF dışa aktarım" in text
    assert D.read_text(b"bozuk", "x.docx") == "" and D.read_text(b"x", "a.png") == ""
    assert D.is_readable("a.docx") and not D.is_readable("a.png") and D.is_unreadable_image("mockup.PNG")


def test_upload_reads_word_files_into_the_notes(client, hub):
    res = client.post("/api/specs/upload", data={"title": "Müşteri Keşif ve Gereksinim Notu", "client_name": "A", "auto_generate_tasks": "false"},
                      files=[("files", ("sartname.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                             ("files", ("mockup.png", b"\x89PNG\r\n", "image/png"))]).json()
    notes = res["spec"]["meeting_notes"]
    assert "Kullanıcılar sisteme kayıt olup giriş yapabilmeli." in notes and res["spec"]["title"] == "Müşteri Portalı"
    assert {f["filename"] for f in res["spec"]["files"]} == {"sartname.docx", "mockup.png"}      # görsel kasada, içeriği okunmadı
